# -*- coding: utf-8 -*-
"""
Convierte una lista de invitados en Excel a un archivo .glf de Guest List.

Cada evento llega con un Excel distinto, porque cada dueno de evento lo arma a su
manera. Lo unico que siempre aparece es la hoja con el plano de mesas, asi que esa
es la fuente principal: cada renglon del plano es una entrada de Guest List.
"Ana Gomez y Pedro Ruiz" es UNA entrada con +1 acompanante, igual que
como se registra a mano en la app.

Formatos de plano que entiende (los tres vistos hasta ahora, y sus mezclas):

    MESA 1  (9 PUESTOS)     MESA 2 (9 PUESTOS)      encabezados en varias columnas,
    ANA GOMEZ Y ...       MARIELA                 los invitados debajo de cada uno
    CARLOS SOTO Y ...     LUISA PARRA 

    Mesa 1 ( 11 puestos )  Adulto  Ninos  Capitan   encabezado y debajo los
    Sr Jorge Medina         2                      invitados con sus cupos
    Sr Hugo Salas           2       1              en columnas aparte

    Mesa #1                                         encabezado sin capacidad y
    Marta Diaz & Julian Ortiz                       "Total: N" al pie
    ...
    Total: 11

Si ademas hay una hoja con la lista de contactos, de ahi salen los telefonos y los
confirmados que todavia no tienen mesa. Si no la hay, se trabaja solo con el plano.
"""

import os
import re
import argparse
import difflib
import unicodedata
import collections

import openpyxl

import glf


class ExcelNoEntendido(Exception):
    """El archivo no tiene una forma que se pueda leer. El mensaje explica que falta."""


# Verdes con que algunas listas marcan a los confirmados. Solo se filtra por color
# cuando la hoja de verdad usa colores; si no, entran todas las filas.
VERDES = {'FF6AA84F', 'FFD9EAD3', 'FF00FF00', 'FFB6D7A8', 'FF93C47D', 'FFC6EFCE'}

# Encabezado -> clave interna. Se busca por subcadena sobre el texto normalizado.
COLUMNAS = [
    ('nombre',     ('invitado', 'nombre', 'familia', 'guest', 'name', 'acompanante')),
    ('cupos',      ('cupo', 'pax', 'puesto', 'persona', 'adulto', 'asiento', 'total')),
    ('ninos',      ('nino', 'ninos', 'child', 'kid')),
    # telefono va antes que contacto a proposito: "Numero de contacto" es un
    # telefono, y si contacto mirara primero se quedaria con esa columna.
    ('telefono',   ('numero', 'telefono', 'celular', 'phone', 'movil', 'whatsapp')),
    ('contacto',   ('contacto', 'capitan')),
    ('email',      ('email', 'correo', 'mail')),
    ('confirmado', ('confirm',)),
    ('mesa',       ('mesa', 'table', 'ubicacion')),
    ('notas',      ('nota', 'observ', 'comentario')),
]

# Palabras que aparecen dentro de un nombre pero no son una persona mas.
RUIDO = {'sr', 'sra', 'srta', 'srita', 'dr', 'dra', 'don', 'dona', 'ing', 'familia',
         'flia', 'pareja', 'hijo', 'hija', 'hijos', 'esposo', 'esposa', 'invitado',
         'acompanante', 'acompanantes', 'invitados', 'su', 'de', 'del', 'la', 'los', 'las',
         'y', 'e', 'n', 'c'}

# Pero si una parte del nombre es solo una de estas, es alguien sin nombre:
# "Samuel y acompanante" son dos personas.
SIN_NOMBRE = {'acompanante', 'invitado', 'pareja', 'esposo', 'esposa', 'hijo', 'hija'}

CABEZA_MESA = re.compile(r'^\s*mesa\s*#?\s*(\d+)\b', re.I)
CABEZA_SIN_NUMERO = re.compile(r'^\s*(novios?|mesa\s+principal|mesa\s+de\s+honor)\b', re.I)
CAPACIDAD = re.compile(r'\(\s*(\d+)\s*[^)]*\)')
CIERRA_BLOQUE = re.compile(r'^\s*(total|totales)\b|^\s*\d+\s*(puestos|personas|cupos)\b', re.I)
MARCA_NINO = re.compile(r'\(\s*(n|ni[nñ][oa]s?|hij[oa]s?|beb[eé]s?)\s*\)', re.I)
PARENTESIS = re.compile(r'\([^)]*\)')
# "Familia Guzman (4)": el parentesis con solo un numero es el cupo de la entrada.
CUPO_ENTRE_PARENTESIS = re.compile(r'\(\s*(\d+)\s*\)')


# ---------------------------------------------------------------- utilidades

def limpiar(valor):
    """Texto de celda sin marcas invisibles y con las 'y' que algun Excel dano como '&'."""
    if valor is None:
        return ''
    s = str(valor)
    for ch in ('‭', '‬', '‎', '‪', '‏', '\xa0'):
        s = s.replace(ch, ' ')
    letra = 'A-Za-zÁÉÍÓÚÜÑáéíóúüñ'
    s = re.sub('(?<=[' + letra + '])&(?=[' + letra + '])', 'y', s)   # "Marí&" raro    
    s = re.sub('(?<=[' + letra + '])&(?![' + letra + '\\s])', 'y', s)
    s = re.sub('(?<![' + letra + '\\s])&(?=[' + letra + '])', 'y', s)
    return re.sub(r'\s{2,}', ' ', s).strip()


def normalizar(texto):
    """Minusculas sin tildes ni signos, como lista de palabras."""
    t = unicodedata.normalize('NFD', texto.lower())
    t = ''.join(c for c in t if unicodedata.category(c) != 'Mn')
    return re.sub(r'[^a-z0-9 ]', ' ', t).split()


def es_numero(texto):
    return bool(re.fullmatch(r'\d+([.,]0+)?', str(texto).strip()))


def entero(valor, por_defecto=0):
    try:
        return int(float(str(valor).strip().replace(',', '.')))
    except (ValueError, TypeError):
        m = re.search(r'\d+', str(valor or ''))
        return int(m.group()) if m else por_defecto


def color(celda):
    relleno = celda.fill
    if not relleno or relleno.patternType != 'solid':
        return None
    return getattr(relleno.fgColor, 'rgb', None)


def personas(texto):
    """Separa 'Sr. X y Sra. Y, Z (nina)' en una lista de personas.

    Devuelve una lista de listas de palabras. Cuantas devuelve es justo el cupo de
    la entrada: comprobado contra las capacidades que los propios Excel declaran.
    """
    t = PARENTESIS.sub(' ', texto)                       # quita "(NIÑA)", "(inhedima egwu)"
    salida = []
    # re.I es indispensable: media lista viene en mayusculas ("X Y MAURICIO LAHORE").
    for parte in re.split(r'\s*[&,;/+]\s*|\s+\by\b\s+|\s+\be\b\s+', t, flags=re.I):
        todas = normalizar(parte)
        palabras = [w for w in todas if w not in RUIDO]
        if len(palabras) == 1 and palabras[0].isdigit():
            # "+2", "y 2 acompanantes": ese numero de personas sin nombre.
            salida.extend([] for _ in range(int(palabras[0])))
        elif palabras:
            salida.append(palabras)
        elif SIN_NOMBRE.intersection(todas):
            # Lista vacia: cuenta en el cupo pero no se cruza con ningun contacto.
            salida.append([])
    return salida


def contar_ninos(texto):
    """Ninos marcados junto al nombre: (n), (NIÑA), (NIÑOS), (HIJO)."""
    marcas = MARCA_NINO.findall(texto)
    if not marcas:
        return 0
    # "(NIÑOS)" en plural marca a todos los de esa entrada.
    if len(marcas) == 1 and marcas[0].lower().endswith('s') and len(personas(texto)) > 1:
        return len(personas(texto))
    return len(marcas)


# ------------------------------------------------- lectura de la hoja de lista

def localizar_encabezado(hoja, max_filas=20):
    """Devuelve (fila, {clave: columna}) o (None, None) si la hoja no parece una lista."""
    mejor = None
    for i, fila in enumerate(hoja.iter_rows(min_row=1, max_row=max_filas, values_only=True), 1):
        textos = [' '.join(normalizar(limpiar(c))) for c in fila]
        mapa = {}
        for clave, palabras in COLUMNAS:
            for j, t in enumerate(textos):
                if t and any(p in t for p in palabras) and j not in mapa.values():
                    mapa.setdefault(clave, j)
                    break
        if 'nombre' in mapa and len(mapa) >= 3 and (mejor is None or len(mapa) > len(mejor[1])):
            mejor = (i, mapa)
    return mejor if mejor else (None, None)


def leer_lista(hoja, mapa, fila_encabezado):
    """Lee las filas de invitados. Devuelve (filas, hay_colores)."""
    filas = []
    col_nombre = mapa['nombre']

    def dato(fila, clave):
        if clave not in mapa or mapa[clave] >= len(fila):
            return ''
        return limpiar(fila[mapa[clave]].value)

    for fila in hoja.iter_rows(min_row=fila_encabezado + 1, max_row=hoja.max_row):
        if col_nombre >= len(fila):
            continue
        celda = fila[col_nombre]
        nombre = limpiar(celda.value)
        if not nombre or es_numero(nombre):
            continue
        cabeza = normalizar(nombre)[:1]
        if cabeza in ([], ['total'], ['totales'], ['lista'], ['mesa']):
            continue
        filas.append({
            'fila': celda.row, 'nombre': nombre, 'color': color(celda),
            'cupos': dato(fila, 'cupos'), 'ninos': dato(fila, 'ninos'),
            'contacto': dato(fila, 'contacto'), 'telefono': dato(fila, 'telefono'),
            'email': dato(fila, 'email'), 'confirmado': dato(fila, 'confirmado'),
            'mesa': dato(fila, 'mesa'), 'notas': dato(fila, 'notas'),
        })
    hay_colores = any(f['color'] in VERDES for f in filas)
    return filas, hay_colores


# ------------------------------------------------------------ plano de mesas

def contar_cabezas_mesa(hoja):
    """Cuantos encabezados de mesa tiene la hoja. Sirve para adivinar cual es el plano."""
    n = 0
    for fila in hoja.iter_rows(values_only=True):
        for v in fila:
            if v is not None and CABEZA_MESA.match(limpiar(v)):
                n += 1
    return n


def leer_plano_mesas(hoja, max_blancos=2):
    """Lee el plano de una hoja de Excel."""
    return leer_plano(
        [[limpiar(c) for c in fila] for fila in hoja.iter_rows(values_only=True)],
        max_blancos)


def leer_plano(grid, max_blancos=2):
    """Lee el plano de una rejilla de celdas ya limpias.

    Va aparte de la hoja de Excel a proposito: una lista pegada como texto se
    vuelve una rejilla de una sola columna y pasa por aqui igual, sin duplicar
    nada de la logica de bloques, capacidades y cupos.

    Devuelve (entradas, capacidades, mesas_repetidas).
    entradas: lista de dicts con mesa (int o None), nombre, cupos y ninos.
    capacidades: {mesa: puestos declarados}, sacado de "(9 PUESTOS)" o de "Total: 9".
    """
    entradas, capacidades = [], {}
    vistas, repetidas = set(), set()

    for i, fila in enumerate(grid):
        for j, valor in enumerate(fila):
            m = CABEZA_MESA.match(valor)
            sin_numero = None if m else CABEZA_SIN_NUMERO.match(valor)
            if not m and not sin_numero:
                continue
            mesa = int(m.group(1)) if m else None
            if mesa is not None and mesa in vistas:
                repetidas.add(mesa)
            vistas.add(mesa)

            # Capacidad declarada, de dos formas: "(9 PUESTOS)" en el propio
            # encabezado, o "Total: 9" mas abajo en la misma columna. Lo segundo se
            # busca aparte de las entradas, porque suele haber filas vacias en medio.
            cap = CAPACIDAD.search(valor)
            if mesa is not None and cap:
                capacidades[mesa] = int(cap.group(1))
            elif mesa is not None:
                for k in range(i + 1, len(grid)):
                    texto = grid[k][j] if j < len(grid[k]) else ''
                    if texto and (CABEZA_MESA.match(texto) or CABEZA_SIN_NUMERO.match(texto)):
                        break
                    if texto and CIERRA_BLOQUE.match(texto):
                        n = re.search(r'\d+', texto)
                        if n:
                            capacidades[mesa] = int(n.group())
                        break

            blancos = 0
            for k in range(i + 1, len(grid)):
                texto = grid[k][j] if j < len(grid[k]) else ''
                if not texto:
                    blancos += 1
                    if blancos >= max_blancos:
                        break
                    continue
                blancos = 0
                if CABEZA_MESA.match(texto) or CABEZA_SIN_NUMERO.match(texto):
                    break
                if CIERRA_BLOQUE.match(texto):
                    break
                if es_numero(texto):
                    continue

                # Columnas a la derecha: si son numeros, son los cupos de esta entrada.
                numeros = []
                for extra in grid[k][j + 1:j + 4]:
                    if es_numero(extra):
                        numeros.append(entero(extra))
                    elif extra:
                        break

                # El cupo, por orden de confianza: las columnas numericas de al
                # lado, un numero entre parentesis, o contar los nombres.
                gente = personas(texto)
                entre_parentesis = CUPO_ENTRE_PARENTESIS.search(texto)
                if numeros:
                    cupos = sum(numeros)
                elif entre_parentesis:
                    cupos = max(1, int(entre_parentesis.group(1)))
                else:
                    cupos = max(1, len(gente))

                entradas.append({
                    'mesa': mesa,
                    'nombre': texto,
                    'cupos': cupos,
                    'ninos': numeros[1] if len(numeros) >= 2 else contar_ninos(texto),
                    'fila': k + 1,
                })
    return entradas, capacidades, sorted(x for x in repetidas if x is not None)


# --------------------------------------------------- listas pegadas en texto

# Basura que traen las listas de WhatsApp y que no es parte del nombre.
SELLO_WHATSAPP = re.compile(r'^\s*\[?\d{1,2}[/-]\d{1,2}[/-]\d{2,4},?\s*'
                            r'\d{1,2}:\d{2}[^\]]*\]?\s*[^:]{0,40}:\s*', re.I)
VINETA = re.compile(r'^\s*(\d{1,3}\s*[.)\-]|[-*•·●–—>+])\s*')
CUPO_AL_FINAL = re.compile(r'\s*[-–—:=]\s*(\d{1,2})\s*$|\s*\bx\s*(\d{1,2})\s*$', re.I)


def limpiar_linea(linea):
    """Un renglon de texto pegado, sin la basura de alrededor.

    Devuelve (nombre, cupos_declarados) donde cupos puede ser None.
    """
    t = limpiar(linea.replace('\t', '  '))
    t = SELLO_WHATSAPP.sub('', t)      # "[12/05, 3:45 p.m.] Ana Gomez: ..."
    t = VINETA.sub('', t)              # "1.", "1)", "-", "*", "•"

    cupos = None
    m = CUPO_AL_FINAL.search(t)
    if m:
        # Solo si delante queda un nombre de verdad; "Mesa 1 - 8" es un encabezado.
        posible = t[:m.start()].strip()
        if posible and not CABEZA_MESA.match(posible) and normalizar(posible):
            cupos = int(m.group(1) or m.group(2))
            t = posible
    return t.strip(), cupos


def leer_plano_texto(texto, max_blancos=2):
    """Lee una lista pegada como texto. Devuelve lo mismo que leer_plano().

    Si el texto trae encabezados de mesa se aprovechan igual que en el Excel; si
    no, cada renglon es una entrada suelta sin mesa.
    """
    lineas, cupos_por_linea = [], {}
    for cruda in texto.splitlines():
        nombre, cupos = limpiar_linea(cruda)
        if cupos is not None:
            cupos_por_linea[len(lineas)] = cupos
        lineas.append(nombre)

    # Una sola columna: el lector de bloques del Excel sirve tal cual.
    grid = [[l] for l in lineas]
    entradas, capacidades, repetidas = leer_plano(grid, max_blancos)

    if not entradas:
        # Sin encabezados de mesa: cada renglon con nombre es un invitado.
        for i, linea in enumerate(lineas):
            if not linea or es_numero(linea) or CIERRA_BLOQUE.match(linea):
                continue
            gente = personas(linea)
            if not gente:
                continue
            entre_parentesis = CUPO_ENTRE_PARENTESIS.search(linea)
            entradas.append({
                'mesa': None,
                'nombre': linea,
                'cupos': (cupos_por_linea.get(i)
                          or (int(entre_parentesis.group(1)) if entre_parentesis else 0)
                          or max(1, len(gente))),
                'ninos': contar_ninos(linea),
                'fila': i + 1,
            })
    else:
        # Los cupos escritos al final del renglon ("Ana y Pedro - 2") mandan.
        for e in entradas:
            declarado = cupos_por_linea.get(e['fila'] - 1)
            if declarado:
                e['cupos'] = declarado
    return entradas, capacidades, repetidas


# ------------------------------------------------ emparejar nombres parecidos

def _iguales(a, b):
    if a == b:
        return True
    if len(a) >= 4 and len(b) >= 4 and (a.startswith(b) or b.startswith(a)):
        return True
    return len(a) >= 5 and len(b) >= 5 and difflib.SequenceMatcher(None, a, b).ratio() >= .8


def _puntaje(p, q):
    """Que tanto se parecen dos personas. 0 = no son la misma."""
    if not p or not q or not _iguales(p[0], q[0]):
        return 0
    ap, aq = set(p[1:]), set(q[1:])
    if not ap or not aq:
        return 3
    comunes = sum(1 for x in ap for y in aq if _iguales(x, y))
    return 3 + comunes * 4 if comunes else 2


def emparejar(grupos_a, grupos_b):
    """Cruza dos colecciones de entradas, cada una con varias personas dentro.

    Cada persona se usa una sola vez, para que dos invitados de nombre parecido
    (dos "Jaime Restrepo") no terminen apuntando al mismo. Una entrada compartida
    ("Gustavo Pardo y Lourdes Vega") si puede casar con dos del otro lado,
    porque cada una casa con una persona distinta.

    Devuelve ({indice_en_a: indice_en_b}, indices_de_b_ya_cubiertos).
    """
    candidatos = []
    for ia, pa in enumerate(grupos_a):
        for ib, pb in enumerate(grupos_b):
            for na, a in enumerate(pa):
                puntos = max((_puntaje(a, b) for b in pb), default=0)
                if puntos > 0:
                    candidatos.append((puntos, ia, na, ib))
    candidatos.sort(reverse=True)

    pares, usados_b, usadas_a = {}, set(), set()
    for _, ia, na, ib in candidatos:
        if ib in usados_b or (ia, na) in usadas_a:
            continue
        usados_b.add(ib)
        usadas_a.add((ia, na))
        pares.setdefault(ia, ib)
    return pares, usados_b


# --------------------------------------------------------------- conversion

def _notas(partes):
    return ' | '.join(p for p in partes if p)


def _contacto_util(contacto, nombre):
    """El nombre del contacto solo aporta si no es alguien ya nombrado en la entrada."""
    if not contacto or es_numero(contacto):
        return ''
    if set(normalizar(contacto)) & set(normalizar(nombre)):
        return ''
    return 'Contacto: %s' % contacto


def _ninos_nota(n):
    return '%d niño%s' % (n, 's' if n > 1 else '') if n else ''


def invitado_desde_plano(entrada, fuente):
    ninos = entrada['ninos']
    return {
        'nombre': entrada['nombre'],
        'apellido': '',
        'telefono': fuente.get('telefono', ''),
        'email': fuente.get('email', ''),
        'mesa': str(entrada['mesa']) if entrada['mesa'] is not None else '',
        'notas': _notas([fuente.get('notas', ''), _ninos_nota(ninos),
                         _contacto_util(fuente.get('contacto', ''), entrada['nombre'])]),
        'acompanantes': max(0, entrada['cupos'] - 1),
        'confirmado': 1,
        'invitado': 1,
    }


def invitado_desde_lista(registro, con_mesa=True):
    gente = personas(registro['nombre'])
    cupos = entero(registro['cupos'], 0) or max(1, len(gente))
    ninos = entero(registro['ninos'], 0) or contar_ninos(registro['nombre'])
    confirmado = 1
    texto_conf = registro['confirmado'].strip().lower()
    if texto_conf:
        confirmado = 0 if texto_conf in ('no', '0', '-', 'x') else 1
        if es_numero(texto_conf):
            confirmado = 1 if entero(texto_conf, 0) > 0 else 0
    return {
        'nombre': registro['nombre'],
        'apellido': '',
        'telefono': registro['telefono'],
        'email': registro['email'],
        'mesa': registro['mesa'] if con_mesa else '',
        'notas': _notas([registro['notas'], _ninos_nota(ninos),
                         _contacto_util(registro['contacto'], registro['nombre'])]),
        'acompanantes': max(0, cupos - 1),
        'confirmado': confirmado,
        'invitado': 1,
    }


# ------------------------------------------------------------------ proceso

SIN_HOJA = '(ninguna)'


def hojas_probables(libro):
    """Adivina que hoja tiene el plano de mesas y cual la lista de contactos."""
    puntajes = {h: contar_cabezas_mesa(libro[h]) for h in libro.sheetnames}
    mesas = max(puntajes, key=puntajes.get) if any(puntajes.values()) else None

    lista = None
    for h in libro.sheetnames:
        if h == mesas:
            continue
        if localizar_encabezado(libro[h])[0]:
            lista = h
            break
    if lista is None and mesas is None:
        lista = libro.sheetnames[0]
    return lista, mesas


def convertir(ruta_excel, hoja=None, hoja_mesas=None, agregar_sin_mesa=False,
              todos=False, amper_y=False):
    """Hace la conversion y devuelve (invitados, informe).

    `hoja` y `hoja_mesas` en None significan "adivinala"; el texto '(ninguna)'
    significa "esa hoja no existe en este archivo".
    """
    libro = openpyxl.load_workbook(ruta_excel, data_only=True)
    auto_lista, auto_mesas = hojas_probables(libro)

    nombre_lista = auto_lista if hoja is None else (None if hoja == SIN_HOJA else hoja)
    nombre_mesas = auto_mesas if hoja_mesas is None else (None if hoja_mesas == SIN_HOJA else hoja_mesas)
    if nombre_lista and nombre_lista == nombre_mesas:
        nombre_lista = None       # una sola hoja: es el plano, no una lista aparte

    inf = {'archivo': os.path.basename(ruta_excel), 'hoja_lista': nombre_lista,
           'hoja_mesas': nombre_mesas, 'avisos': [], 'sin_mesa': [],
           'mesas_que_no_cuadran': [], 'columnas': {}, 'verdes': 0, 'descartados': 0,
           'descartados_por_color': collections.Counter(), 'filtro_color': False}

    # --- hoja de contactos (opcional)
    lista = []
    if nombre_lista:
        ws = libro[nombre_lista]
        fila_enc, mapa = localizar_encabezado(ws)
        if fila_enc is None:
            inf['avisos'].append(
                'La hoja "%s" no tiene una fila de encabezados reconocible '
                '(esperaba algo como "Invitados", "Nombre" o "Familia"); la ignore.'
                % nombre_lista)
            inf['hoja_lista'] = None
            nombre_lista = None
        else:
            todas, hay_colores = leer_lista(ws, mapa, fila_enc)
            inf['fila_encabezado'] = fila_enc
            inf['columnas'] = {k: (chr(65 + v) if v < 26 else str(v + 1)) for k, v in mapa.items()}
            if hay_colores and not todos:
                lista = [f for f in todas if f['color'] in VERDES]
                fuera = [f for f in todas if f['color'] not in VERDES]
                inf['filtro_color'] = True
                inf['descartados'] = len(fuera)
                inf['descartados_por_color'] = collections.Counter(f['color'] for f in fuera)
            else:
                lista = todas
                if not hay_colores:
                    inf['avisos'].append(
                        'La hoja "%s" no usa colores, asi que tome todas sus filas.' % nombre_lista)
            inf['verdes'] = len(lista)

    # --- plano de mesas (fuente principal)
    if nombre_mesas:
        entradas, capacidades, repetidas = leer_plano_mesas(libro[nombre_mesas])
        if not entradas:
            raise ExcelNoEntendido(
                'En la hoja "%s" no encontre ninguna mesa. Espero celdas que empiecen '
                'por "Mesa 1", "MESA 1 (9 PUESTOS)" o "Mesa #1", con los invitados '
                'escritos debajo en la misma columna.' % nombre_mesas)
        invitados = armar_desde_plano(entradas, capacidades, repetidas, lista, inf,
                                      agregar_sin_mesa)
    else:
        if not lista:
            raise ExcelNoEntendido(
                'Este archivo no tiene ni plano de mesas ni una lista con encabezados '
                'que pueda reconocer. Revisa que la hoja elegida sea la correcta.')
        invitados = [invitado_desde_lista(f) for f in lista]
        inf['fuente'] = 'lista'
        inf['avisos'].append('No encontre plano de mesas; use la hoja de la lista.')

    return rematar(invitados, inf, amper_y)


def armar_desde_plano(entradas, capacidades, repetidas, lista, inf, agregar_sin_mesa):
    """Convierte las entradas del plano en invitados y llena el informe.

    Lo usan por igual el Excel y el texto pegado, para que el resultado y las
    comprobaciones sean exactamente los mismos vengan de donde vengan.
    """
    if repetidas:
        inf['avisos'].append(
            'El plano trae el encabezado de la mesa %s mas de una vez. Junte todo '
            'lo que aparece debajo de cada uno; revisa que no haya un bloque viejo '
            'olvidado.' % ', '.join(str(x) for x in repetidas))

    pares, cubiertos = emparejar([personas(e['nombre']) for e in entradas],
                                 [personas(f['nombre']) for f in lista])
    invitados = [invitado_desde_plano(e, lista[pares[i]] if i in pares else {})
                 for i, e in enumerate(entradas)]

    inf['fuente'] = 'plano'
    inf['n_mesas'] = len(set(e['mesa'] for e in entradas if e['mesa'] is not None))
    inf['n_entradas_plano'] = len(entradas)
    inf['personas_plano'] = sum(1 + g['acompanantes'] for g in invitados)
    inf['sin_numero_de_mesa'] = [e['nombre'] for e in entradas if e['mesa'] is None]

    contados = collections.Counter()
    for e in entradas:
        if e['mesa'] is not None:
            contados[e['mesa']] += e['cupos']
    inf['mesas_revisadas'] = len(capacidades)
    inf['mesas_que_no_cuadran'] = [(m, capacidades[m], contados[m])
                                   for m in sorted(capacidades)
                                   if capacidades[m] != contados[m]]
    if not capacidades and inf['n_mesas']:
        inf['avisos'].append(
            'El plano no declara cuantos puestos tiene cada mesa, asi que no pude '
            'contrastar los cupos. Vale la pena revisar la lista a ojo.')

    # El plano es lo ultimo que arman los novios, asi que manda sobre la hoja de
    # contactos. Que alguien este en la lista y no en el plano suele querer decir
    # que lo sacaron, no que le falta mesa: por eso no entra salvo que se pida.
    pendientes = [f for i, f in enumerate(lista) if i not in cubiertos]
    inf['sin_mesa'] = [(f['fila'], f['nombre']) for f in pendientes]
    inf['sin_mesa_agregados'] = bool(pendientes) and agregar_sin_mesa
    if inf['sin_mesa_agregados']:
        for f in pendientes:
            g = invitado_desde_lista(f, con_mesa=False)
            g['notas'] = _notas([g['notas'], 'No esta en el plano de mesas'])
            invitados.append(g)
    return invitados


def rematar(invitados, inf, amper_y=False):
    """Los ultimos retoques y las cuentas finales del informe."""
    if amper_y:
        for g in invitados:
            g['nombre'] = re.sub(r'\s*&\s*', ' y ', g['nombre'])
    invitados.sort(key=lambda g: normalizar(g['nombre']))

    inf['sin_telefono'] = [g['nombre'] for g in invitados if not g['telefono']]
    inf['entradas'] = len(invitados)
    inf['personas'] = sum(1 + g['acompanantes'] for g in invitados)
    return invitados, inf


def convertir_texto(texto, agregar_sin_mesa=False, amper_y=False):
    """Convierte una lista pegada como texto, tipica de WhatsApp.

    Devuelve lo mismo que convertir(): (invitados, informe).
    """
    entradas, capacidades, repetidas = leer_plano_texto(texto)
    inf = {'archivo': 'Lista pegada como texto', 'hoja_lista': None,
           'hoja_mesas': None, 'avisos': [], 'sin_mesa': [],
           'mesas_que_no_cuadran': [], 'columnas': {}, 'verdes': 0, 'descartados': 0,
           'descartados_por_color': collections.Counter(), 'filtro_color': False,
           'es_texto': True,
           'lineas': len([l for l in texto.splitlines() if l.strip()])}

    if not entradas:
        raise ExcelNoEntendido(
            'No encontre ningun invitado en ese texto. Pega la lista con un nombre '
            'por renglon; si trae mesas, deja los renglones "Mesa 1", "Mesa 2"... '
            'con sus invitados debajo.')

    invitados = armar_desde_plano(entradas, capacidades, repetidas, [], inf,
                                  agregar_sin_mesa)
    if not inf['n_mesas']:
        inf['avisos'].append(
            'El texto no traia mesas, asi que todos los invitados quedaron sin mesa. '
            'Se las puedes poner despues en Guest List.')
    return rematar(invitados, inf, amper_y)


def resumen_fuente(inf):
    """Una frase que dice de donde salio la lista. Va arriba del resultado."""
    if inf.get('es_texto'):
        con_mesas = ('con sus mesas' if inf.get('n_mesas')
                     else 'sin mesas: el texto no las traia')
        return 'Lei %d renglones de la lista que pegaste, %s.' % (inf['lineas'], con_mesas)
    if inf['fuente'] == 'lista':
        return 'Sin plano de mesas: lei los invitados de la hoja "%s".' % inf['hoja_lista']
    if inf['hoja_lista']:
        return ('Lei el plano de mesas de "%s" y los telefonos de "%s".'
                % (inf['hoja_mesas'], inf['hoja_lista']))
    return 'Lei el plano de mesas de "%s". Este archivo no trae hoja de contactos.' % inf['hoja_mesas']


def hallazgos(inf):
    """Lo que hay que revisar, como lista de cosas concretas en vez de un muro de texto.

    Cada hallazgo es {nivel, titulo, detalle}. nivel: 'ojo' (revisar), 'info' (para
    saber) u 'ok' (salio bien). La ventana los pinta como tarjetas y la consola los
    imprime en lineas.
    """
    H = []
    # La lista puede venir de un Excel o de un texto pegado: el informe tiene que
    # nombrar bien de donde salieron los puestos declarados.
    origen = 'la lista' if inf.get('es_texto') else 'el Excel'

    if inf['fuente'] == 'plano':
        if inf['mesas_que_no_cuadran']:
            H.append({
                'nivel': 'ojo',
                'titulo': '%d mesa%s no cuadra%s con lo que dice %s'
                          % (len(inf['mesas_que_no_cuadran']),
                             '' if len(inf['mesas_que_no_cuadran']) == 1 else 's',
                             '' if len(inf['mesas_que_no_cuadran']) == 1 else 'n',
                             origen),
                'detalle': ['Mesa %d: dice %d puestos, yo cuento %d personas'
                            % (m, d, c) for m, d, c in inf['mesas_que_no_cuadran']]})
        elif inf['mesas_revisadas']:
            n = inf['mesas_revisadas']
            H.append({'nivel': 'ok',
                      'titulo': ('La mesa cuadra con los puestos que declara %s' % origen
                                 if n == 1 else
                                 'Las %d mesas cuadran con los puestos que declara %s'
                                 % (n, origen)),
                      'detalle': []})

        if inf['sin_mesa']:
            nombres = ['%s   (fila %d)' % (n, f) for f, n in inf['sin_mesa']]
            if inf['sin_mesa_agregados']:
                H.append({'nivel': 'info',
                          'titulo': '%d de la hoja de contactos no estan en el plano, '
                                    'pero los agregaste sin mesa' % len(nombres),
                          'detalle': nombres})
            else:
                H.append({'nivel': 'ojo',
                          'titulo': '%d de la hoja de contactos no estan en el plano, '
                                    'asi que no van en la lista' % len(nombres),
                          'detalle': nombres + [
                              '',
                              'El plano es lo ultimo que arman los novios, asi que lo mas '
                              'probable es que ya no vayan.',
                              'Si alguno si va, marca la casilla de abajo y vuelve a leer.']})

        # Si no hay mesas en absoluto no vale la pena listarlos a todos: eso ya lo
        # dice el aviso de arriba.
        if inf.get('sin_numero_de_mesa') and inf.get('n_mesas'):
            H.append({'nivel': 'info',
                      'titulo': '%d entrada%s de un bloque sin numero de mesa'
                                % (len(inf['sin_numero_de_mesa']),
                                   '' if len(inf['sin_numero_de_mesa']) == 1 else 's'),
                      'detalle': inf['sin_numero_de_mesa'] +
                                 ['', 'Van en la lista, pero sin mesa asignada.']})

    if inf['sin_telefono']:
        if len(inf['sin_telefono']) == inf['entradas']:
            H.append({'nivel': 'info', 'titulo': 'Nadie quedo con telefono',
                      'detalle': ['Una lista pegada como texto no trae telefonos.'
                                  if inf.get('es_texto') else
                                  'Este Excel no trae una columna de telefonos.']})
        else:
            H.append({'nivel': 'info',
                      'titulo': '%d invitados sin telefono' % len(inf['sin_telefono']),
                      'detalle': inf['sin_telefono']})

    if inf['filtro_color'] and inf['descartados']:
        detalle = ['%d con relleno %s' % (n, c or 'sin color')
                   for c, n in inf['descartados_por_color'].most_common()]
        H.append({'nivel': 'info',
                  'titulo': '%d filas de la hoja de contactos descartadas por su color'
                            % inf['descartados'],
                  'detalle': detalle + ['', 'Solo entran las filas verdes.']})

    for a in inf['avisos']:
        H.append({'nivel': 'info', 'titulo': a, 'detalle': []})

    orden = {'ojo': 0, 'info': 1, 'ok': 2}
    H.sort(key=lambda h: orden[h['nivel']])
    if not any(h['nivel'] == 'ojo' for h in H):
        H.sort(key=lambda h: 0 if h['nivel'] == 'ok' else 1)
    return H


def texto_informe(inf):
    """El mismo informe de la ventana, en texto, para la consola."""
    MARCA = {'ojo': '[ OJO ]', 'info': '[  i  ]', 'ok': '[  OK ]'}
    L = ['Archivo: %s' % inf['archivo'],
         resumen_fuente(inf),
         '',
         '%d invitados  |  %d personas%s'
         % (inf['entradas'], inf['personas'],
            '  |  %d mesas' % inf['n_mesas'] if inf.get('n_mesas') else ''),
         '']
    encontrados = hallazgos(inf)
    if not encontrados:
        L.append('Todo en orden, no hay nada que revisar.')
    for h in encontrados:
        L.append('%s %s' % (MARCA[h['nivel']], h['titulo']))
        for d in h['detalle']:
            L.append('        %s' % d if d else '')
        L.append('')
    return '\n'.join(L).rstrip()


# ---------------------------------------------------------------------- main

def main(argv=None):
    p = argparse.ArgumentParser(description='Convierte un Excel de invitados a .glf de Guest List.')
    p.add_argument('excel', help='Ruta del archivo .xlsx')
    p.add_argument('--hoja', help='Hoja con la lista de contactos ("%s" si no hay)' % SIN_HOJA)
    p.add_argument('--mesas', help='Hoja con el plano de mesas ("%s" si no hay)' % SIN_HOJA)
    p.add_argument('--evento', help='Nombre del evento = nombre del archivo .glf')
    p.add_argument('--fecha', help='Fecha del evento, formato 2026-08-29')
    p.add_argument('--hora', default='18:00', help='Hora del evento (por defecto 18:00)')
    p.add_argument('--salida', help='Ruta del .glf de salida')
    p.add_argument('--agregar-sin-mesa', action='store_true',
                   help='Agregar tambien los de la hoja de contactos que no estan en el plano')
    p.add_argument('--todos', action='store_true', help='No filtrar por color')
    p.add_argument('--y', dest='amper_y', action='store_true',
                   help='Escribir " y " en vez de " & " en los nombres')
    args = p.parse_args(argv)

    try:
        invitados, inf = convertir(args.excel, hoja=args.hoja, hoja_mesas=args.mesas,
                                   agregar_sin_mesa=args.agregar_sin_mesa,
                                   todos=args.todos, amper_y=args.amper_y)
    except ExcelNoEntendido as e:
        raise SystemExit('No pude leer el archivo:\n%s' % e)
    print(texto_informe(inf))

    evento = args.evento or os.path.splitext(os.path.basename(args.excel))[0]
    salida = args.salida or os.path.join(os.path.dirname(os.path.abspath(args.excel)), evento + '.glf')
    glf.escribir_glf(salida, invitados, fecha=args.fecha, hora=args.hora)

    print('')
    print('Listo: %s' % salida)
    print('%d entradas en la lista, %d personas en total.' % (inf['entradas'], inf['personas']))
    return salida


if __name__ == '__main__':
    main()
